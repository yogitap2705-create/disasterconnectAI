import os
import shutil
import subprocess

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from report_content_part1 import add_preliminary_and_chapters_1_to_3_part1
from report_content_part2 import add_chapter_3_part2
from report_content_part3 import add_chapter_3_part3
from report_content_part4 import add_chapter_3_part4_and_remaining_chapters

ASSETS_DIR = r"C:\Users\ammyp\OneDrive\Desktop\DisasterConnectAI\report_assets"
OUTPUT_DOCX = r"C:\Users\ammyp\OneDrive\Desktop\DisasterConnectAI\DisasterConnectAI_Project_Report.docx"
DOWNLOADS_DOCX = r"C:\Users\ammyp\Downloads\DisasterConnectAI_Project_Report.docx"
ROOT_DOCX = r"C:\DisasterConnectAI\DisasterConnectAI_Project_Report.docx"

print("Initializing Master Document Assembly...")
doc = docx.Document()

# 1. Margins (Left margin 1.5 inches for punch / binding space as requested)
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)

# 2. Double Page Border (Sant Rawool Maharaj Mahavidyalaya Format)
def add_page_border(section):
    sectPr = section._sectPr
    pgBorders = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'  <w:top w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:left w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:bottom w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
        f'  <w:right w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(pgBorders)

add_page_border(doc.sections[0])

# 3. Typography Helpers
def p_center(text, size=12, bold=False, italic=False, space_before=0, space_after=6, color=(0,0,0), underline=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.color.rgb = RGBColor(*color)
    return p

def p_heading1(text, space_before=16, space_after=10, color=(128, 0, 32)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.underline = True
    run.font.color.rgb = RGBColor(*color)
    return p

def p_heading2(text, space_before=12, space_after=6, color=(26, 54, 93)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*color)
    return p

def p_heading3(text, space_before=8, space_after=4, color=(45, 55, 72)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*color)
    return p

def p_body(text, space_before=0, space_after=6, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.italic = italic
    return p

def p_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p

def p_code(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 30, 30)
    return p

def add_image(filename, width_inch=5.5, caption=""):
    full_path = os.path.join(ASSETS_DIR, filename)
    if os.path.exists(full_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(full_path, width=Inches(width_inch))
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r = p_cap.add_run(caption)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.italic = True
            r.font.bold = True
            r.font.color.rgb = RGBColor(60, 60, 60)
    else:
        print(f"Warning: Image {filename} not found in {ASSETS_DIR}")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="B0B0B0", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_table(headers, data_rows, col_widths, bg_header="1A365D", align_center_cols=None):
    if align_center_cols is None:
        align_center_cols = [0]
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CBD5E0", sz="4")

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], bg_header)

    # Data Rows
    for r_idx, row in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in align_center_cols else WD_ALIGN_PARAGRAPH.LEFT
            if len(p.runs) > 0:
                p.runs[0].font.name = "Times New Roman"
                p.runs[0].font.size = Pt(8.5)
                p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
            set_cell_background(row_cells[c_idx], bg_col)

    # Column Widths (Auto-scaled to fit within printable width: 5.75 inches)
    tot_w = sum(col_widths)
    max_avail_w = 5.75
    scale = (max_avail_w / tot_w) if tot_w > max_avail_w else 1.0
    scaled_widths = [w * scale for w in col_widths]
    for row in table.rows:
        for i, w in enumerate(scaled_widths):
            row.cells[i].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

helpers = {
    'p_center': p_center,
    'p_heading1': p_heading1,
    'p_heading2': p_heading2,
    'p_heading3': p_heading3,
    'p_body': p_body,
    'p_bullet': p_bullet,
    'p_code': p_code,
    'add_image': add_image,
    'add_styled_table': add_styled_table
}

# Assemble all parts in order
print("Adding Preliminary Pages, Chapter I, Chapter II, and Chapter III (Sections 1-5)...")
add_preliminary_and_chapters_1_to_3_part1(doc, helpers)

print("Adding Chapter III (Sections 6-16: Architecture, ERD, Data Dictionary, UML, Math)...")
add_chapter_3_part2(doc, helpers)

print("Adding Chapter III (Sections 17-20: 8 Windows Desktop Screens, SOPs, 40 Test Cases)...")
add_chapter_3_part3(doc, helpers)

print("Adding Chapter III (Section 21), Chapter IV (Annexures A-D), Chapters V, VI, VII...")
add_chapter_3_part4_and_remaining_chapters(doc, helpers)

# Save master document
print(f"Saving Master Word Report to {OUTPUT_DOCX}...")
doc.save(OUTPUT_DOCX)
print("Saved successfully!")

# Copy to Downloads and C: drive
shutil.copyfile(OUTPUT_DOCX, DOWNLOADS_DOCX)
shutil.copyfile(OUTPUT_DOCX, ROOT_DOCX)
print("Copied successfully to Desktop, Downloads, and C:\\DisasterConnectAI!")

# Verify exact page count using Word COM
print("Checking exact Word page count via COM...")
try:
    ps_cmd = f"""
    $w = New-Object -ComObject Word.Application
    $w.Visible = $false
    $d = $w.Documents.Open('{OUTPUT_DOCX}')
    $pages = $d.ComputeStatistics(2)
    Write-Host "OFFICIAL_WORD_PAGE_COUNT: $pages"
    $d.Close([ref]$false)
    $w.Quit()
    """
    res = subprocess.run(["powershell", "-Command", ps_cmd], capture_output=True, text=True)
    print(res.stdout)
except Exception as e:
    print("Page count check error:", e)

print("Assembly and Verification Complete!")
